"""Draws the 100-ticket validation sample and writes validation/validation_sheet.xlsx for HAND labelling.
The sheet deliberately hides the rule-based label so the human label is independent.
Usage (from repo root, after data/ is populated):  python python/validation_sample.py
"""
import re, pathlib, duckdb, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

root = pathlib.Path(__file__).resolve().parents[1]
con = duckdb.connect()
for stmt in re.split(r";\s*(?:\n|$)", (root / "sql" / "01_clean.sql").read_text()):
    if re.sub(r"--[^\n]*", "", stmt).strip():
        con.execute(stmt)

# Stratified sample of closed tickets that HAVE closing text: 15 per bucket for six buckets, 10 from 'other'. Fixed ordering = reproducible.
quota = {"no_finding": 15, "action_taken": 15, "could_not_reach_or_gone": 15, "info_or_referral": 15,
         "duplicate": 15, "fixed_per_caller_unverified": 15, "other": 10}
parts = []
for bucket, n in quota.items():
    parts.append(con.execute(f"""SELECT unique_key, agency, complaint_type, resolution_description AS closing_note
        FROM tickets WHERE status = 'Closed' AND outcome = '{bucket}'
        ORDER BY hash(unique_key || 'validation-v1') LIMIT {n}""").fetchdf())
sample = pd.concat(parts).sample(frac=1, random_state=7).reset_index(drop=True)
sample.insert(0, "n", range(1, len(sample) + 1))
sample["your_label"] = ""
assert len(sample) == 100

LABELS = [
    ("action_taken", "The agency did something about it: fixed it, cleaned it, repaired it, issued a summons or violation, did the requested work."),
    ("no_finding", "The agency checked and found nothing: no violation, no evidence, no condition, or said no action was necessary."),
    ("could_not_reach_or_gone", "The agency could not complete the job: no access, nobody home, the person or vehicle was gone on arrival, or could not find the person."),
    ("duplicate", "Closed because it duplicates another complaint or an earlier complaint at the same location."),
    ("fixed_per_caller_unverified", "Closed because someone on the phone (caller, tenant or occupant) said the problem was already fixed. No inspection result."),
    ("info_or_referral", "Information only, a pointer to a website, a referral to another agency, or a note that the request was reviewed or fulfilled with no inspection result."),
    ("other", "None of the above fits."),
]

wb = Workbook(); ws = wb.active; ws.title = "Label"
ws.append(["n", "unique_key", "agency", "complaint_type", "closing_note (read this)", "your_label (pick from the dropdown)"])
for r in sample.itertuples(index=False):
    ws.append([r.n, r.unique_key, r.agency, r.complaint_type, r.closing_note, ""])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F"); c.alignment = Alignment(wrap_text=True, vertical="center")
for col, w in zip("ABCDEF", [5, 13, 9, 26, 95, 34]): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    row[5].fill = PatternFill("solid", fgColor="FFF4E5")
dv = DataValidation(type="list", formula1='"' + ",".join(l for l, _ in LABELS) + '"', allow_blank=False)
ws.add_data_validation(dv); dv.add(f"F2:F{ws.max_row}")
ws.freeze_panes = "A2"

g = wb.create_sheet("Guide")
g.append(["Label", "Choose it when..."])
for l, d in LABELS: g.append([l, d])
g.append([]); g.append(["Rules for you", "Read only the closing note. Label what the note says happened, not what you think happened. Do not look at any rule output. One label per row."])
g.column_dimensions["A"].width = 30; g.column_dimensions["B"].width = 120
for row in g.iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
g["A1"].font = g["B1"].font = Font(bold=True)

out = root / "validation"; out.mkdir(exist_ok=True)
wb.save(out / "validation_sheet.xlsx")
print("wrote", out / "validation_sheet.xlsx", "rows:", len(sample))
